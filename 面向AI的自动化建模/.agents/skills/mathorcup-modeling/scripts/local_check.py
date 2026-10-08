"""Bounded, offline capability checks. Nothing is installed or reconfigured."""
from __future__ import annotations

import argparse
import csv
import importlib
import importlib.metadata
import importlib.util
import json
import math
import os
from pathlib import Path
import platform
import shutil
import subprocess
import sys
from datetime import datetime, timezone

FEATURES = {
    "base": (), "numerics": ("numpy", "scipy", "pandas"), "learning": ("sklearn",),
    "plot": ("matplotlib", "seaborn"), "excel": ("openpyxl",), "pdf": ("pymupdf",),
    "latex": (),
}
DISTRIBUTIONS = {"pymupdf": "pymupdf", "sklearn": "scikit-learn"}
DEFAULT = "numerics,plot,excel,pdf,latex"


def parse_features(value):
    selected = list(dict.fromkeys(x.strip() for x in value.split(",") if x.strip()))
    unknown = set(selected) - FEATURES.keys()
    if not selected or unknown:
        raise ValueError("Choose features from: " + ", ".join(FEATURES))
    return selected


def module_info(name):
    try:
        actual = "fitz" if name == "pymupdf" and importlib.util.find_spec("pymupdf") is None else name
        module = importlib.import_module(actual)
        try:
            version = importlib.metadata.version(DISTRIBUTIONS.get(name, name))
        except importlib.metadata.PackageNotFoundError:
            version = getattr(module, "__version__", "unknown")
        return {"available": True, "version": version or "unknown", "import_tested": True}
    except Exception as exc:
        return {"available": False, "reason": str(exc)}


def find_xelatex():
    candidates = [os.environ.get("MATHORCUP_XELATEX"), shutil.which("xelatex")]
    appdata = os.environ.get("APPDATA")
    if appdata:
        candidates.append(str(Path(appdata) / "TinyTeX/bin/windows/xelatex.exe"))
    for root in (Path.home() / ".TinyTeX/bin", Path.home() / "Library/TinyTeX/bin"):
        if root.is_dir():
            candidates.extend(str(p) for p in root.glob("*/xelatex"))
    for value in candidates:
        if value and Path(value).is_file():
            return str(Path(value).resolve())
    return None


def probe(selected):
    modules = {name: module_info(name) for f in selected for name in FEATURES[f]}
    report = {
        "checked_utc": datetime.now(timezone.utc).isoformat(),
        "python": sys.version.split()[0], "executable": sys.executable,
        "platform": platform.system(), "mode": "import_check",
        "features": {}, "modules": modules,
    }
    version_ok = sys.version_info >= (3, 10)
    for feature in selected:
        missing = [m for m in FEATURES[feature] if not modules[m]["available"]]
        item = {"available": version_ok and not missing, "unavailable_modules": missing}
        if feature == "latex":
            engine = find_xelatex()
            item.update(available=version_ok and bool(engine), executable=engine)
        report["features"][feature] = item
    report["ok"] = version_ok and all(v["available"] for v in report["features"].values())
    report["note"] = "Imports were checked, not model execution. Access denied may be an execution policy issue, not a missing package. Use smoke for end-to-end checks."
    return report


def sample_rows():
    return [(i, 2 * i + 1) for i in range(6)]


def numerics(out):
    import numpy as np
    import pandas as pd
    from scipy.optimize import Bounds, LinearConstraint, linprog, milp

    frame = pd.read_csv(out / "synthetic_data.csv")
    if frame.to_numpy().tolist() != [list(row) for row in sample_rows()]:
        raise AssertionError("Pandas CSV parsing changed the synthetic data")

    a = np.array([[1.0, 1.0], [2.0, 1.0]])
    b = np.array([4.0, 5.0])
    c = np.array([-3.0, -2.0])
    lp = linprog(c, A_ub=a, b_ub=b, bounds=(0, None), method="highs")
    mip = milp(c, integrality=np.ones(2), bounds=Bounds(0, np.inf),
               constraints=LinearConstraint(a, -np.inf, b), options={"time_limit": 10})
    if not lp.success or not mip.success:
        raise RuntimeError(f"Solver status: LP={lp.message}; MILP={mip.message}")
    # A separate scalar enumeration checks this deliberately tiny integer case.
    exact = max(3*x + 2*y for x in range(6) for y in range(6)
                if x+y <= 4 and 2*x+y <= 5)
    residual = max(0.0, float(np.max(a @ lp.x - b)))
    if residual > 1e-8 or not math.isclose(-lp.fun, exact, abs_tol=1e-8):
        raise AssertionError("LP failed the independent feasibility/objective check")
    if not math.isclose(-mip.fun, exact, abs_tol=1e-8):
        raise AssertionError("MILP differs from enumeration")
    result = {"synthetic_environment_example": True, "lp_x": lp.x.tolist(),
              "lp_objective": float(-lp.fun), "milp_objective": float(-mip.fun),
              "enumerated_objective": exact, "max_violation": residual}
    (out / "synthetic_optimization.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
    return result


def learning(out):
    import numpy as np
    from sklearn.linear_model import LinearRegression
    rows = np.array(sample_rows(), dtype=float)
    model = LinearRegression().fit(rows[:4, :1], rows[:4, 1])
    error = float(np.max(np.abs(model.predict(rows[4:, :1]) - rows[4:, 1])))
    if error > 1e-8:
        raise AssertionError("Synthetic regression smoke check failed")
    return {"synthetic_environment_example": True, "max_prediction_error": error,
            "note": "A library smoke test, not evidence about contest model quality."}


def plot(out):
    import warnings
    import matplotlib
    matplotlib.use("Agg")
    import seaborn as sns
    from matplotlib import pyplot as plt, font_manager, ft2font
    text = "环境自检（合成数据）时段数值"
    families = ("Microsoft YaHei", "SimHei", "Noto Sans CJK SC", "Source Han Sans SC", "SimSun", "PingFang SC")
    font = None
    for family in families:
        try:
            path = font_manager.findfont(font_manager.FontProperties(family=family), fallback_to_default=False)
            face = ft2font.FT2Font(path)
            if all(face.get_char_index(ord(ch)) for ch in text):
                font = path
                break
        except (ValueError, OSError, RuntimeError):
            continue
    if font is None:
        raise RuntimeError("No verified Chinese font found; choose an existing CJK font. No font was downloaded.")
    prop = font_manager.FontProperties(fname=font)
    with matplotlib.rc_context({"axes.unicode_minus": False, "pdf.fonttype": 42}):
        fig, ax = plt.subplots(figsize=(6.4, 3.5), layout="constrained")
        try:
            rows = sample_rows()
            sns.lineplot(x=[r[0] for r in rows], y=[r[1] for r in rows], marker="o", ax=ax)
            ax.set_title("环境自检（合成数据）", fontproperties=prop)
            ax.set_xlabel("时段", fontproperties=prop)
            ax.set_ylabel("数值", fontproperties=prop)
            with warnings.catch_warnings(record=True) as caught:
                warnings.simplefilter("always")
                for suffix in ("png", "pdf", "svg"):
                    fig.savefig(out / f"synthetic_figure.{suffix}", dpi=160)
            glyph_errors = [str(w.message) for w in caught if "Glyph" in str(w.message) and "missing" in str(w.message)]
            if glyph_errors:
                raise RuntimeError("; ".join(glyph_errors))
        finally:
            plt.close(fig)
    return {"font": Path(font).name, "formats": ["png", "pdf", "svg"], "synthetic_environment_example": True}


def excel(out):
    from openpyxl import Workbook, load_workbook
    target = out / "synthetic_table.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "环境合成样例"
    sheet.append(["x", "y"])
    for row in sample_rows():
        sheet.append(row)
    book.save(target)
    book.close()
    reopened = load_workbook(target, read_only=True, data_only=True)
    try:
        actual = list(reopened.active.values)[1:]
        if actual != sample_rows():
            raise AssertionError("Excel round trip changed values")
    finally:
        reopened.close()
    return {"round_trip_rows": len(actual), "formula_recalculation_tested": False}


def pdf(out):
    fitz = importlib.import_module("pymupdf" if importlib.util.find_spec("pymupdf") else "fitz")
    source = out / "synthetic_figure.pdf"
    if not source.is_file():
        source = out / "synthetic_pdf.pdf"
        with fitz.open() as doc:
            doc.new_page().insert_text((72, 72), "SYNTHETIC ENVIRONMENT CHECK: 1 3 5 7 9 11")
            doc.save(source)
    with fitz.open(source) as doc:
        if len(doc) != 1 or not doc[0].get_text().strip():
            raise AssertionError("PDF page/text check failed")
        pix = doc[0].get_pixmap(matrix=fitz.Matrix(1.3, 1.3))
        pix.save(out / "pdf_render.png")
        return {"pages": len(doc), "render_pixels": [pix.width, pix.height],
                "note": "Programmatic checks passed; inspect the rendered page for final visual QA."}


def latex(out):
    engine = find_xelatex()
    version = subprocess.run([engine, "--version"], capture_output=True, text=True,
                             encoding="utf-8", errors="replace", timeout=10)
    if version.returncode != 0:
        raise RuntimeError("XeLaTeX version check failed")
    if "miktex" in (version.stdout + version.stderr).lower():
        raise RuntimeError("MiKTeX auto-install policy is unknown; use an existing no-install configuration.")
    source = out / "synthetic_latex.tex"
    source.write_text(r"\documentclass[UTF8,fontset=fandol]{ctexart}" + "\n" +
                      r"\usepackage{amsmath}\begin{document}" + "\n" +
                      r"\section*{环境自检：合成样例}" + "\n" +
                      r"这里只核验中文、公式和实际编译，不是比赛结果。" + "\n" +
                      r"\begin{equation}\label{eq:check}x+y\leq 4,\qquad 2x+y\leq 5.\end{equation}" + "\n" +
                      r"交叉引用检查：式~\eqref{eq:check}。" + "\n" +
                      r"\end{document}" + "\n", encoding="utf-8")
    target = out / "synthetic_latex.pdf"
    logs = []
    for pass_number in (1, 2):
        run = subprocess.run([engine, "-no-shell-escape", "-interaction=nonstopmode", "-halt-on-error", source.name],
                             cwd=out, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=40)
        log = run.stdout + run.stderr
        logs.append(f"PASS {pass_number}\n{log}")
        (out / "latex_console.txt").write_text("\n".join(logs), encoding="utf-8")
        if run.returncode != 0 or not target.is_file():
            raise RuntimeError(f"LaTeX pass {pass_number} failed; see latex_console.txt. No packages were installed.")
        if "Missing character:" in log:
            raise RuntimeError("LaTeX reported missing glyphs")
    if "undefined references" in log.lower() or "Rerun to get cross-references right" in log:
        raise RuntimeError("LaTeX cross-references did not resolve after two passes")
    return {"engine": engine, "exit_code": run.returncode, "pdf_bytes": target.stat().st_size,
            "version": version.stdout.splitlines()[0], "passes": 2, "cross_references_resolved": True}


RUNNERS = {"base": lambda out: {"stdlib": True}, "numerics": numerics, "learning": learning,
           "plot": plot, "excel": excel, "pdf": pdf, "latex": latex}


def smoke(selected, target):
    report = probe(selected)
    report["mode"] = "synthetic_smoke"
    if not report["ok"]:
        return report, 2
    out = Path(target).expanduser().resolve()
    # Refuse all existing directories: never touch original data or prior results.
    out.mkdir(parents=True, exist_ok=False)
    report["output"] = str(out)
    report["synthetic_environment_example"] = True
    report["checks"] = {}
    with (out / "synthetic_data.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(["synthetic_x", "synthetic_y"])
        writer.writerows(sample_rows())
    for feature in FEATURES:
        if feature not in selected:
            continue
        try:
            report["checks"][feature] = {"ok": True, "details": RUNNERS[feature](out)}
        except Exception as exc:
            report["checks"][feature] = {"ok": False, "error": f"{type(exc).__name__}: {exc}"}
    report["ok"] = all(item["ok"] for item in report["checks"].values())
    report["note"] = "Only these synthetic capability checks ran; this is not validation of a contest solution."
    (out / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report, 0 if report["ok"] else 1


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("mode", choices=["check", "smoke"])
    parser.add_argument("--features", default=DEFAULT)
    parser.add_argument("--out", help="New directory for explicitly synthetic smoke artifacts")
    parser.add_argument("--quiet", action="store_true", help="Suppress report; retain exit status")
    options = parser.parse_args(argv)
    try:
        selected = parse_features(options.features)
        if options.mode == "smoke":
            if not options.out:
                raise ValueError("smoke requires --out with a new directory")
            report, code = smoke(selected, options.out)
        else:
            report = probe(selected)
            code = 0 if report["ok"] else 2
        if not options.quiet:
            print(json.dumps(report, ensure_ascii=False, indent=2))
        return code
    except (ValueError, OSError) as exc:
        if not options.quiet:
            print(json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False))
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
